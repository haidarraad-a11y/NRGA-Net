from docx import Document

d = Document(r"D:\Dataset\pp3\NRGA-Net_paper_revised.docx")

# check no placeholders remain anywhere
bad = []
def scan(text, where):
    for marker in ("TO BE FILLED", "[P]", "[R]", "[F1]", "[IoU]", "= F1]",
                   "97.35", "99.77", "1.6843", "2 / 525", "27 / 525", "0.05%",
                   "placeholder", "surpassing FECDNet by 1.92"):
        if marker in text:
            bad.append((where, marker, text[:90]))

for i, p in enumerate(d.paragraphs):
    scan(p.text, f"para {i}")
for ti, t in enumerate(d.tables):
    for r in t.rows:
        for c in r.cells:
            scan(c.text, f"table {ti}")

print("leftover markers:", len(bad))
for b in bad[:20]:
    print("  ", b)

# verify key numbers present
full = "\n".join(p.text for p in d.paragraphs)
for key in ("97.80", "99.22", "98.89", "96.58", "98.71", "T = 0.50", "4 / 3,076",
            "20,811", "16,966", "3,076", "99.71"):
    print(f"{key!r}: {'OK' if key in full else 'MISSING'}")

# spot-check tables
for ti in (3, 4, 5, 6):
    t = d.tables[ti]
    print(f"== table {ti} ({len(t.rows)}r)")
    for r in t.rows[:4]:
        print("   " + " | ".join(c.text.strip()[:24] for c in r.cells))
