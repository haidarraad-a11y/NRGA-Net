from docx import Document

d = Document(r"D:\Dataset\pp3\NRGA-Net_paper.docx")
for i, t in enumerate(d.tables):
    first = t.cell(0, 0).text.strip()[:30]
    print(f"== table {i}: {len(t.rows)}r x {len(t.columns)}c | first: {first}")
    if i in (3, 4, 5):
        for r in t.rows[:12]:
            print("   " + " | ".join(c.text.strip()[:26] for c in r.cells))
